"""Automated tests for FairHire helper logic.

Run with either:
    python -m unittest -v test_fairhire
    pytest -v test_fairhire.py
"""

import io
import json
import sqlite3
import sys
import tempfile
import os
import unittest

import app  # must NOT launch the Streamlit interface on import
import auth


def make_pdf(pages):
    """Build a small text PDF, one string per page (no external libraries)."""

    def esc(line):
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", f"<< /Type /Pages /Kids [{kids}] /Count {count} >>", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = "BT /F1 11 Tf 14 TL 50 740 Td " + " ".join(f"({esc(line)}) Tj T*" for line in text.split("\n")) + " ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


R1_P1 = "Priya Sharma\npriya.sharma@example.com | +91 98765 43210\nSummary\n5 years of experience building machine learning models.\nSkills: Python, SQL, pandas"
R1_P2 = "Experience\n2019 - 2024 Data Scientist at Example Analytics\nBuilt dashboards with Tableau"
R2_P1 = "Rahul Verma\nrahul.v@example.org\nExperience\nBackend developer, 3 years of experience in Java and SQL"
R3_P1 = "Sneha Iyer\nSenior ML engineer\nSkills\nPython, PyTorch, NLP, deep learning, AWS\n8 years of experience"


class ImportSafetyTests(unittest.TestCase):
    def test_import_does_not_load_streamlit(self):
        self.assertNotIn("streamlit", sys.modules, "importing app must not start/import Streamlit")

    def test_main_is_defined_but_not_called(self):
        self.assertTrue(callable(app.main))


class SkillExtractionTests(unittest.TestCase):
    def test_basic_skills(self):
        found = app.extract_skills("Experienced with Python, SQL and Machine Learning.")
        self.assertEqual(found, ["machine learning", "python", "sql"])

    def test_special_character_skills(self):
        found = app.extract_skills("I write C++ and C# and Node.js.")
        self.assertIn("c++", found)
        self.assertIn("c#", found)
        self.assertIn("node.js", found)

    def test_no_partial_word_matches(self):
        found = app.extract_skills("I know JavaScript and MySQL.")
        self.assertNotIn("java", found)
        self.assertNotIn("sql", found)
        self.assertIn("javascript", found)
        self.assertIn("mysql", found)

    def test_aliases_and_case(self):
        self.assertIn("scikit-learn", app.extract_skills("used SKLEARN daily"))

    def test_empty_inputs(self):
        self.assertEqual(app.extract_skills(""), [])
        self.assertEqual(app.extract_skills(None), [])
        self.assertFalse(app.has_skill("", "python"))
        self.assertFalse(app.has_skill("python", ""))


class ExperienceTests(unittest.TestCase):
    def test_stated_years(self):
        self.assertEqual(app.extract_experience_years("5 years of experience in data"), 5.0)
        self.assertEqual(app.extract_experience_years("8+ years of professional experience"), 8.0)

    def test_date_ranges_merged(self):
        text = "2015 - 2018 Dev at A\n2017 - 2020 Dev at B"
        self.assertEqual(app.extract_experience_years(text, current_year=2024), 5.0)

    def test_present_uses_current_year(self):
        self.assertEqual(app.extract_experience_years("2020 - Present Engineer", current_year=2024), 4.0)

    def test_education_dates_ignored(self):
        text = "2014 - 2018 B.Sc University of X"
        self.assertEqual(app.extract_experience_years(text, current_year=2024), 0.0)

    def test_empty_and_capped(self):
        self.assertEqual(app.extract_experience_years(""), 0.0)
        self.assertLessEqual(app.extract_experience_years("99 years of experience"), 40.0)


class AnonymizationTests(unittest.TestCase):
    SAMPLE = (
        "Jane Doe\n"
        "Email: jane.doe@example.com | Phone: +1 (415) 555-2671\n"
        "https://www.linkedin.com/in/janedoe\n"
        "Gender: Female\n"
        "She worked at Acme from 2018 - 2022.\n"
    )

    def setUp(self):
        self.out = app.anonymize_text(self.SAMPLE)

    def test_email_phone_url_removed(self):
        self.assertNotIn("jane.doe@example.com", self.out)
        self.assertNotIn("555-2671", self.out)
        self.assertNotIn("linkedin.com", self.out)

    def test_name_and_gender_removed(self):
        self.assertNotIn("Jane Doe", self.out)
        self.assertNotIn("Female", self.out)
        self.assertNotIn("She ", self.out)

    def test_years_not_treated_as_phone(self):
        self.assertIn("2018 - 2022", self.out)

    def test_heading_first_line_not_masked(self):
        self.assertTrue(app.anonymize_text("Work Experience\nPython developer").startswith("Work Experience"))

    def test_skill_or_title_first_line_is_not_mistaken_for_a_name(self):
        self.assertTrue(app.anonymize_text("Python SQL\nsome text").startswith("Python SQL"))
        self.assertTrue(app.anonymize_text("Python Developer\nsome text").startswith("Python Developer"))
        self.assertTrue(app.anonymize_text("Jane Doe\nsome text").startswith("[NAME]"))

    def test_empty(self):
        self.assertEqual(app.anonymize_text(""), "")
        self.assertEqual(app.anonymize_text(None), "")


class ScoringTests(unittest.TestCase):
    JD = "Data scientist using Python and SQL to build machine learning models."

    def test_bounds(self):
        for text in ["", "python", "python sql machine learning 20 years of experience", "nothing relevant"]:
            result = app.score_candidate(text, ["python", "sql"], self.JD, 3)
            self.assertGreaterEqual(result["score"], 0.0)
            self.assertLessEqual(result["score"], 100.0)

    def test_perfect_and_zero(self):
        perfect = app.score_candidate("Python SQL. 10 years of experience.", ["python", "sql"], "", 5)
        self.assertEqual(perfect["score"], 100.0)
        zero = app.score_candidate("Cooking and gardening", ["python", "sql"], "", 5)
        self.assertEqual(zero["score"], 0.0)

    def test_deterministic(self):
        a = app.score_candidate("Python and SQL, 4 years of experience", ["python", "sql", "aws"], self.JD, 3)
        b = app.score_candidate("Python and SQL, 4 years of experience", ["python", "sql", "aws"], self.JD, 3)
        self.assertEqual(a, b)

    def test_explanations(self):
        result = app.score_candidate("Python developer", ["python", "sql"], "", 0)
        self.assertEqual(result["matched_skills"], ["python"])
        self.assertEqual(result["missing_skills"], ["sql"])
        self.assertTrue(any("Missing" in r for r in result["reasons"]))
        self.assertEqual(result["score"], 50.0)

    def test_empty_resume(self):
        result = app.score_candidate("   ", ["python"], self.JD, 2)
        self.assertEqual(result["score"], 0.0)
        self.assertEqual(result["missing_skills"], ["python"])

    def test_empty_job_and_no_skills(self):
        result = app.score_candidate("Python developer", [], "", 0)
        self.assertEqual(result["score"], 0.0)
        self.assertTrue(result["reasons"])

    def test_skills_derived_from_job_description(self):
        result = app.score_candidate("Python and SQL expert", [], self.JD, 0)
        self.assertIn("python", result["matched_skills"])
        self.assertGreater(result["score"], 0.0)

    def test_more_matches_never_lower_score(self):
        low = app.score_candidate("Python", ["python", "sql", "aws"], "", 0)["score"]
        high = app.score_candidate("Python SQL", ["python", "sql", "aws"], "", 0)["score"]
        self.assertGreater(high, low)

    def test_group_labels_have_no_effect(self):
        base = "Python and SQL, 4 years of experience"
        a = app.score_candidate(base + " Group A", ["python", "sql"], "", 3)["score"]
        b = app.score_candidate(base + " Group B", ["python", "sql"], "", 3)["score"]
        self.assertEqual(a, b)


class ScreeningPipelineTests(unittest.TestCase):
    def test_ranking_and_duplicates(self):
        docs = [
            ("a.txt", "Python SQL 5 years of experience"),
            ("b.txt", "Python only"),
            ("dup.txt", "python   only"),  # duplicate of b after normalisation
            ("empty.txt", "   "),
        ]
        results, warnings = app.screen_resumes(docs, ["python", "sql"], "", 2, anonymize=True)
        self.assertEqual([r["file_name"] for r in results], ["a.txt", "b.txt"])
        self.assertEqual([r["rank"] for r in results], [1, 2])
        self.assertEqual(len(warnings), 2)

    def test_no_documents(self):
        results, warnings = app.screen_resumes([], ["python"], "", 0)
        self.assertEqual(results, [])
        self.assertEqual(warnings, [])

    def test_dataframe_and_csv(self):
        results, _ = app.screen_resumes([("x.txt", "Python")], ["python"], "", 0)
        df = app.results_to_dataframe(results)
        self.assertNotIn("File", df.columns)
        self.assertIn("File", app.results_to_dataframe(results, include_file_names=True).columns)
        self.assertIsInstance(app.dataframe_to_csv_bytes(df), bytes)
        self.assertTrue(app.results_to_dataframe([]).empty)

    def test_csv_export_neutralises_formulas(self):
        import pandas as pd

        df = pd.DataFrame({"File": ["=HYPERLINK(1)", "ok.txt"], "Score": [1.0, 2.0]})
        data = app.dataframe_to_csv_bytes(df).decode("utf-8")
        self.assertIn("'=HYPERLINK(1)", data)
        self.assertIn("ok.txt", data)

    def test_csv_injection_guard(self):
        self.assertEqual(app.csv_safe("=SUM(A1)"), "'=SUM(A1)")
        self.assertEqual(app.csv_safe("normal"), "normal")


class FileReadingTests(unittest.TestCase):
    def test_txt(self):
        text, err = app.read_resume_bytes("cv.txt", b"Hello Python")
        self.assertEqual(text, "Hello Python")
        self.assertIsNone(err)

    def test_unsupported_empty_and_corrupt(self):
        self.assertIsNotNone(app.read_resume_bytes("cv.exe", b"abc")[1])
        self.assertIsNotNone(app.read_resume_bytes("cv.txt", b"")[1])
        self.assertIsNotNone(app.read_resume_bytes("cv.pdf", b"not a pdf")[1])
        self.assertIsNotNone(app.read_resume_bytes("cv.docx", b"not a docx")[1])

    def test_too_large(self):
        big = b"a" * (app.MAX_FILE_BYTES + 1)
        self.assertIsNotNone(app.read_resume_bytes("cv.txt", big)[1])

    def test_docx_roundtrip(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        document = docx.Document()
        document.add_paragraph("Python developer with SQL")
        buffer = io.BytesIO()
        document.save(buffer)
        text, err = app.read_resume_bytes("cv.docx", buffer.getvalue())
        self.assertIsNone(err)
        self.assertIn("Python developer", text)

    def test_pdf_roundtrip(self):
        try:
            from pypdf import PdfWriter
        except ImportError:
            self.skipTest("pypdf not installed")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        text, err = app.read_resume_bytes("blank.pdf", buffer.getvalue())
        self.assertIsNone(text)  # blank page: no extractable text
        self.assertIn("no extractable text", err)


class FairnessMetricTests(unittest.TestCase):
    def test_rates_and_metrics(self):
        groups = ["A"] * 10 + ["B"] * 10
        selected = [True] * 5 + [False] * 5 + [True] * 2 + [False] * 8
        rates = app.selection_rates(groups, selected)
        self.assertAlmostEqual(rates["A"], 0.5)
        self.assertAlmostEqual(rates["B"], 0.2)
        self.assertAlmostEqual(app.demographic_parity_difference(rates), 0.3)
        self.assertAlmostEqual(app.disparate_impact_ratio(rates), 0.4)

    def test_parity(self):
        rates = app.selection_rates(["A", "A", "B", "B"], [True, False, True, False])
        self.assertEqual(app.demographic_parity_difference(rates), 0.0)
        self.assertEqual(app.disparate_impact_ratio(rates), 1.0)

    def test_edge_cases(self):
        self.assertEqual(app.demographic_parity_difference({"A": 0.5}), 0.0)
        self.assertIsNone(app.disparate_impact_ratio({"A": 0.5}))
        self.assertIsNone(app.disparate_impact_ratio({"A": 0.0, "B": 0.0}))
        self.assertEqual(app.selection_rates([], []), {})
        with self.assertRaises(ValueError):
            app.selection_rates(["A"], [])

    def test_report_flags(self):
        groups = ["A"] * 20 + ["B"] * 20
        selected = [True] * 10 + [False] * 10 + [True] * 2 + [False] * 18
        report = app.fairness_report(groups, selected)
        self.assertIs(report["four_fifths_pass"], False)
        small = app.fairness_report(["A", "B"], [True, False])
        self.assertTrue(any("Small groups" in n for n in small["notes"]))

    def test_select_top_k(self):
        self.assertEqual(app.select_top_k([1, 3, 2], 2), [False, True, True])
        self.assertEqual(app.select_top_k([1, 2], 0), [False, False])
        self.assertEqual(app.select_top_k([1, 2], 99), [True, True])

    def test_audit_uploaded_outcomes(self):
        import pandas as pd

        df = pd.DataFrame({"g": ["A", "A", "B", "B"], "s": ["yes", "no", "yes", "yes"]})
        report = app.audit_uploaded_outcomes(df, "g", "s")
        self.assertAlmostEqual(report["rates"]["A"], 0.5)
        self.assertAlmostEqual(report["rates"]["B"], 1.0)
        with self.assertRaises(ValueError):
            app.audit_uploaded_outcomes(df, "g", "missing")


class SyntheticDataTests(unittest.TestCase):
    def test_reproducible(self):
        a = app.generate_synthetic_candidates(n=100, seed=1)
        b = app.generate_synthetic_candidates(n=100, seed=1)
        self.assertTrue(a.equals(b))

    def test_shape_and_bounds(self):
        df = app.generate_synthetic_candidates(n=120, seed=3, shortlist_fraction=0.25)
        self.assertEqual(len(df), 120)
        self.assertTrue(df["merit_score"].between(0, 100).all())
        self.assertTrue(df["biased_score"].between(0, 100).all())
        self.assertEqual(int(df["selected_merit_only"].sum()), 30)

    def test_zero_candidates(self):
        self.assertEqual(len(app.generate_synthetic_candidates(n=0)), 0)

    def test_simulated_bias_is_detected(self):
        df = app.generate_synthetic_candidates(n=600, seed=7, bias_penalty=30, shortlist_fraction=0.25)
        fair = app.fairness_report(df["group"].tolist(), df["selected_merit_only"].tolist())
        biased = app.fairness_report(df["group"].tolist(), df["selected_biased_process"].tolist())
        self.assertGreater(biased["demographic_parity_difference"], fair["demographic_parity_difference"])
        self.assertLess(biased["disparate_impact_ratio"], 0.8)

    def test_group_not_used_in_merit_score(self):
        df = app.generate_synthetic_candidates(n=200, seed=5, bias_penalty=0)
        self.assertTrue((df["merit_score"] == df["biased_score"]).all())


class UiComponentTests(unittest.TestCase):
    def setUp(self):
        self.results, _ = app.screen_resumes(
            [("<script>alert(1)</script>.txt", "Python SQL. 5 years of experience. $100 & <b>bold</b>"), ("b.txt", "Python")],
            ["python", "sql"], "", 2, anonymize=False,
        )

    def test_theme_css_has_both_palettes_and_fallback(self):
        light, dark = app.theme_css(False), app.theme_css(True)
        for css in (light, dark):
            self.assertTrue(css.startswith("<style>") and css.endswith("</style>"))
            self.assertNotIn("\n", css)
            self.assertIn(":root:has(.fh-dark-marker)", css)
            self.assertIn(":root:not(:has(.fh-dark-marker))", css)
            self.assertIn(app.THEMES["light"]["bg"], css)
            self.assertIn(app.THEMES["dark"]["bg"], css)
        # the plain :root fallback block (first --bg occurrence) follows the requested mode
        self.assertTrue(light.split("--bg: ", 1)[1].startswith(app.THEMES["light"]["bg"]))
        self.assertTrue(dark.split("--bg: ", 1)[1].startswith(app.THEMES["dark"]["bg"]))

    def test_theme_marker(self):
        self.assertIn("fh-dark-marker", app.theme_marker_html(True))
        self.assertNotIn("fh-dark-marker", app.theme_marker_html(False))
        self.assertNotIn("\n", app.theme_marker_html(True))

    def test_theme_script_and_label(self):
        for dark in (False, True):
            script = app.theme_script_html(dark)
            self.assertTrue(script.startswith("<script>") and script.endswith("</script>"))
            self.assertEqual(script.count("</script>"), 1)
            self.assertIn("window.parent.document", script)
            self.assertIn(f"'data-fh-theme','{'dark' if dark else 'light'}'", script)
            self.assertIn(app.THEMES["dark"]["bg"], script)
        self.assertIn("Dark", app.theme_label_html(True))
        self.assertIn("Light", app.theme_label_html(False))
        self.assertIn('data-fh-theme="dark"', app.theme_css(False))

    def test_content_clears_streamlit_fixed_header(self):
        # Streamlit's header is fixed at the top; too little top padding puts the
        # dark-mode toggle underneath it, where clicks never reach the toggle.
        import re

        for rule in re.findall(r"\.block-container \{[^}]*padding: ([\d.]+)rem", app.BASE_CSS):
            self.assertGreaterEqual(float(rule), 4.0)

    def test_intro_is_calm_and_never_blocks_clicks(self):
        import re

        html_out = app.intro_html()
        self.assertIn("Welcome to FairHire", html_out)
        self.assertNotIn("\n", html_out)
        self.assertNotIn("<script", html_out)
        rule = re.search(r"\.fh-intro \{([^}]*)\}", app.BASE_CSS).group(1)
        self.assertIn("pointer-events: none", rule)
        self.assertIn("forwards", rule)  # stays hidden after fading out
        self.assertIn("visibility: hidden", app.BASE_CSS)
        self.assertIn("prefers-reduced-motion", app.BASE_CSS)

    def test_html_is_escaped(self):
        table = app.results_table_html(self.results, 1, show_files=True)
        card = app.candidate_card_html(self.results[0], show_files=True)
        for html_out in (table, card):
            self.assertNotIn("<script>", html_out)
        self.assertIn("&lt;script&gt;", table)
        self.assertNotIn("<b>bold</b>", card)
        self.assertNotIn("$", card)

    def test_no_blank_lines_that_break_markdown(self):
        outputs = [
            app.results_table_html(self.results, 1),
            app.candidate_card_html(self.results[0], open_card=True),
            app.hero_html(),
            app.kpis_html([app.kpi_card("A", 1, "hint", "ok")]),
            app.rates_bars_html({"A": 0.5, "B": 0.25}, {"A": 4, "B": 4}),
            app.empty_state_html("t", "b"),
            app.note_html("line1\n\nline2", "warn", "!"),
        ]
        for out in outputs:
            self.assertNotIn("\n", out)

    def test_candidate_card_open_attribute(self):
        self.assertIn("<details class=\"fh-cand\" open>", app.candidate_card_html(self.results[0], open_card=True))
        self.assertNotIn(" open>", app.candidate_card_html(self.results[0], open_card=False))

    def test_fairness_panel_variants(self):
        bad = app.fairness_report(["A"] * 20 + ["B"] * 20, [True] * 10 + [False] * 10 + [True] * 2 + [False] * 18)
        good = app.fairness_report(["A", "A", "B", "B"], [True, False, True, False])
        undefined = app.fairness_report(["A", "B"], [False, False])
        single = app.fairness_report(["A", "A"], [True, False])
        self.assertIn("Flag for review", app.fairness_panel_html("t", bad))
        self.assertIn("Passes four-fifths guide", app.fairness_panel_html("t", good))
        self.assertIn("Cannot assess", app.fairness_panel_html("t", undefined))
        self.assertIn("Cannot assess", app.fairness_panel_html("t", single))

    def test_score_tone_and_bar_bounds(self):
        self.assertEqual([app.score_tone(x) for x in (0, 49.9, 50, 74.9, 75, 100)], ["bad", "bad", "warn", "warn", "ok", "ok"])
        self.assertIn("width:100.0%", app._score_cell(250))
        self.assertIn("width:0.0%", app._score_cell(-5))

    def test_df_table_html(self):
        import pandas as pd

        df = pd.DataFrame({"a": ["<x>", "y"], "ok": [True, False], "v": [1.234, 2.0]})
        out = app.df_table_html(df, max_rows=1)
        self.assertIn("&lt;x&gt;", out)
        self.assertIn("✓", out)
        self.assertNotIn(">y<", out)
        self.assertNotIn("\n", out)


class CustomSkillAndWeightTests(unittest.TestCase):
    def test_parse_custom_skills(self):
        skills, warnings = app.parse_custom_skills("RAG = retrieval augmented generation, rag pipeline\n# comment\n\nfigma\n=bad")
        self.assertEqual(skills["rag"], ["rag", "retrieval augmented generation", "rag pipeline"])
        self.assertEqual(skills["figma"], ["figma"])
        self.assertEqual(len(warnings), 1)

    def test_custom_skill_is_scored_without_mutating_global_bank(self):
        custom, _ = app.parse_custom_skills("rag = retrieval augmented generation")
        bank = app.merged_skill_bank(custom)
        self.assertNotIn("rag", app.SKILL_BANK)
        result = app.score_candidate("Built Retrieval Augmented Generation systems", ["rag"], "", 0, skill_bank=bank)
        self.assertEqual(result["matched_skills"], ["rag"])
        self.assertEqual(app.canonical_skill("retrieval augmented generation", bank), "rag")

    def test_weights_change_scores_and_stay_bounded(self):
        text = "Python developer, 1 years of experience"
        skills_heavy = app.score_candidate(text, ["python", "sql"], "", 5, weights=(100, 0, 0))["score"]
        exp_heavy = app.score_candidate(text, ["python", "sql"], "", 5, weights=(0, 100, 0))["score"]
        self.assertEqual(skills_heavy, 50.0)
        self.assertEqual(exp_heavy, 20.0)
        for w in [(1, 1, 1), (0, 0, 5), (1000, 0, 0)]:
            self.assertTrue(0 <= app.score_candidate(text, ["python"], "python", 2, weights=w)["score"] <= 100)

    def test_bad_weights_fall_back_to_defaults(self):
        base = app.score_candidate("Python", ["python"], "", 0)["score"]
        for bad in [None, (0, 0, 0), (1, 2), ("a", "b", "c"), (-5, -1, -1)]:
            self.assertEqual(app.score_candidate("Python", ["python"], "", 0, weights=bad)["score"], base)
        self.assertEqual(app.weights_percent((60, 25, 15)), (60.0, 25.0, 15.0))
        self.assertAlmostEqual(sum(app.weights_percent((1, 1, 1))), 100.0, places=0)

    def test_reranking_with_weights(self):
        docs = [("a.txt", "Python SQL"), ("b.txt", "Python SQL 10 years of experience")]
        by_skills, _ = app.screen_resumes(docs, ["python", "sql"], "", 5, weights=(100, 0, 0))
        by_experience, _ = app.screen_resumes(docs, ["python", "sql"], "", 5, weights=(0, 100, 0))
        self.assertEqual(by_experience[0]["file_name"], "b.txt")
        self.assertEqual(by_skills[0]["score"], by_skills[1]["score"])


class ProxyAndRedactionTests(unittest.TestCase):
    def test_flags(self):
        text = (
            "Studied at Example University.\nI am a Hindu temple volunteer. Married with kids.\n"
            "Graduated in 1998. Based in Gurugram.\n2010 - 2012 Dev at A\n2016 - 2019 Dev at B\n"
        )
        cats = {f["category"] for f in app.detect_proxy_risks(text, current_year=2024)}
        for expected in ["Institution", "Religion / community", "Family / marital status", "Age / graduation year", "Location", "Career gap or leave"]:
            self.assertIn(expected, cats)

    def test_no_flags_for_plain_skills_text(self):
        self.assertEqual(app.detect_proxy_risks("Python and SQL developer. 2019 - 2022 Engineer at Acme.", 2024), [])
        self.assertEqual(app.detect_proxy_risks(""), [])

    def test_flags_capped_per_category(self):
        text = " ".join(f"University {i}" for i in range(10))
        flags = [f for f in app.detect_proxy_risks(text) if f["category"] == "Institution"]
        self.assertLessEqual(len(flags), 3)

    def test_results_include_flags_and_original(self):
        results, _ = app.screen_resumes([("x.txt", "Jane Doe\nGraduated 1999.\nPython")], ["python"], "", 0)
        self.assertTrue(results[0]["proxy_flags"])
        self.assertIn("Jane Doe", results[0]["original_text"])
        self.assertNotIn("Jane Doe", results[0]["text_used"])

    def test_redaction_diff_and_count(self):
        original = "Jane Doe\nEmail: jane@example.com\nPython"
        masked = app.anonymize_text(original)
        out = app.redaction_diff_html(original, masked)
        self.assertIn('class="fh-del"', out)
        self.assertIn('class="fh-add"', out)
        self.assertNotIn("\n", out)
        self.assertEqual(app.count_redactions(masked), 2)
        self.assertEqual(app.count_redactions(""), 0)
        self.assertNotIn("<script>", app.redaction_diff_html("<script>x</script>", "y"))


class StressTestTests(unittest.TestCase):
    RESUME = "Priya Sharma\nShe has 4 years of experience with Python and SQL. Her manager praised her work.\n"

    def test_helpers(self):
        self.assertEqual(app.detect_name(self.RESUME), "Priya Sharma")
        self.assertEqual(app.detect_name("Name: Rahul Verma\nPython"), "Rahul Verma")
        self.assertIsNone(app.detect_name("Work Experience\nPython"))
        self.assertEqual(app.swap_gender_terms("He met HER. She"), "She met HIS. He")
        swapped = app.swap_name("Priya Sharma wrote code. Sharma left.", "Priya Sharma", "Emily Carter")
        self.assertNotIn("Sharma", swapped)

    def test_identity_blind_scorer_passes(self):
        for anonymize in (False, True):
            result = app.bias_stress_test(self.RESUME, ["python", "sql"], "", 3, anonymize=anonymize)
            self.assertTrue(result["passed"])
            self.assertEqual(result["max_abs_delta"], 0.0)
            self.assertTrue(result["name_detected"])
            self.assertGreaterEqual(len(result["rows"]), 6)

    def test_detects_a_name_that_leaks_into_scoring(self):
        # the candidate's name happens to be a job-description keyword, so changing it changes the score
        leaky = "Dashboard Rao\nWrites SQL daily."
        result = app.bias_stress_test(leaky, [], "dashboard reporting metrics", 0)
        self.assertTrue(result["name_detected"])
        self.assertFalse(result["passed"])
        self.assertGreater(result["max_abs_delta"], 0)

    def test_without_name_only_gender_test(self):
        result = app.bias_stress_test("he writes python", ["python"], "", 0)
        self.assertFalse(result["name_detected"])
        self.assertEqual(len(result["rows"]), 2)

    def test_stress_html(self):
        results, _ = app.screen_resumes([("a.txt", self.RESUME)], ["python"], "", 0, anonymize=False)
        tests = [app.bias_stress_test(results[0]["original_text"], ["python"], "", 0)]
        self.assertIn("Unchanged", app.stress_summary_html(results, tests))
        detail = app.stress_detail_html(tests[0])
        self.assertIn("Gender terms swapped", detail)
        self.assertNotIn("\n", detail)


class StatisticsTests(unittest.TestCase):
    def test_wilson_interval(self):
        low, high = app.wilson_interval(5, 10)
        self.assertTrue(0 <= low < 0.5 < high <= 1)
        self.assertEqual(app.wilson_interval(0, 0), (0.0, 1.0))
        self.assertEqual(app.wilson_interval(0, 10)[0], 0.0)

    def test_two_proportion_z_test(self):
        r = app.two_proportion_z_test(50, 100, 30, 100)
        self.assertAlmostEqual(r["z"], 2.887, places=2)
        self.assertLess(r["p_value"], 0.01)
        self.assertLess(r["ci_low"], r["diff"] < r["ci_high"] and r["diff"] or 0)
        same = app.two_proportion_z_test(10, 20, 10, 20)
        self.assertEqual(same["p_value"], 1.0)
        self.assertEqual(app.two_proportion_z_test(0, 20, 0, 20)["p_value"], 1.0)
        self.assertIsNone(app.two_proportion_z_test(1, 0, 1, 5))

    def test_report_has_significance_and_intervals(self):
        groups = ["A"] * 100 + ["B"] * 100
        selected = [True] * 50 + [False] * 50 + [True] * 30 + [False] * 70
        report = app.fairness_report(groups, selected)
        self.assertTrue(report["significance"]["significant"])
        self.assertEqual(report["significance"]["high_group"], "A")
        self.assertEqual(report["selected_counts"], {"A": 50, "B": 30})
        for low, high in report["intervals"].values():
            self.assertTrue(0 <= low <= high <= 1)
        small = app.fairness_report(["A", "A", "B", "B"], [True, False, False, False])
        self.assertFalse(small["significance"]["significant"])
        self.assertIsNone(app.fairness_report(["A", "A"], [True, False])["significance"])
        three = app.fairness_report(["A", "B", "C"] * 10, [True, False, False] * 10)
        self.assertTrue(any("multiple-comparison" in n for n in three["notes"]))

    def test_panel_shows_significance(self):
        groups = ["A"] * 100 + ["B"] * 100
        selected = [True] * 50 + [False] * 50 + [True] * 30 + [False] * 70
        out = app.fairness_panel_html("t", app.fairness_report(groups, selected))
        self.assertIn("Statistically significant", out)
        self.assertIn("95% CI", out)


class ReviewAndReportTests(unittest.TestCase):
    def test_validate_review_entry(self):
        self.assertIsNone(app.validate_review_entry("Advance", ""))
        self.assertIsNone(app.validate_review_entry("Reject", "Missing required SQL skills"))
        self.assertIsNotNone(app.validate_review_entry("Reject", "  "))
        self.assertIsNotNone(app.validate_review_entry("Hold", ""))
        self.assertIsNotNone(app.validate_review_entry("Maybe", "x"))
        self.assertIsNotNone(app.validate_review_entry("Advance", "x" * 501))

    def test_review_log_dataframe(self):
        df = app.review_log_dataframe({"k": {"time": "t", "candidate": "Candidate 001", "score": 80.0, "decision": "Advance", "reason": ""}})
        self.assertEqual(list(df.columns), ["Time (UTC)", "Candidate", "Score", "Decision", "Reason"])
        self.assertTrue(app.review_log_dataframe({}).empty)

    def test_histogram(self):
        out = app.score_histogram_html([5, 15, 15, 95, 100, 250, -5])
        self.assertEqual(out.count('class="fh-hist-col"'), 10)
        self.assertNotIn("\n", out)
        self.assertIn("height:103px", out)  # tallest bin
        self.assertTrue(app.score_histogram_html([]).startswith('<div class="fh-hist">'))

    def test_report_html(self):
        results, _ = app.screen_resumes(
            [("<b>x</b>.txt", "Python SQL 5 years of experience. Graduated 1999."), ("y.txt", "Python")], ["python", "sql"], "", 2,
        )
        syn = app.generate_synthetic_candidates(n=100, seed=2, bias_penalty=25)
        sections = app.synthetic_fairness_sections(syn)
        log = [{"time": "now", "candidate": "Candidate 001", "score": 90.0, "decision": "Reject", "reason": "<script>alert(1)</script>"}]
        settings = {"required": ["python", "sql"], "min_years": 2, "weights": (60, 25, 15), "anonymize": True, "shortlist": 1}
        html_out = app.build_report_html(results, settings, sections, log, generated_on="2026-01-01 00:00 UTC", show_files=True)
        self.assertTrue(html_out.startswith("<!DOCTYPE html>"))
        for needle in ["Settings", "Ranking", "Score explanations", "Fairness audit", "Reviewer decisions", "SIMULATED", "Limitations"]:
            self.assertIn(needle, html_out)
        self.assertNotIn("<script>", html_out)
        self.assertNotIn("<b>x</b>", html_out)
        self.assertIn("2026-01-01", html_out)
        empty = app.build_report_html([], {}, None, None)
        self.assertIn("No fairness audit", empty)
        self.assertIn("No reviewer decisions", empty)


class AccountUiHelperTests(unittest.TestCase):
    def test_brand_html(self):
        out = app.auth_brand_html()
        self.assertIn("Smarter hiring starts here", out)
        self.assertNotIn("\n", out)

    def test_job_settings_round_trip(self):
        state = {
            "req_skills": ["python", "sql"], "extra_skills_text": "figma", "custom_skills_text": "figma",
            "min_years": 3.0, "w_skill": 50, "w_exp": 30, "w_kw": 20, "shortlist": 5,
        }
        settings = app.job_settings_from_state(state)
        self.assertNotIn("resume", json.dumps(settings).lower())
        fresh = {}
        app.apply_job_to_state(fresh, {"jd_text": "We need Python", "settings": json.loads(json.dumps(settings))})
        self.assertEqual(fresh["req_skills"], ["python", "sql"])
        self.assertEqual((fresh["w_skill"], fresh["w_exp"], fresh["w_kw"]), (50, 30, 20))
        self.assertEqual((fresh["min_years"], fresh["shortlist"], fresh["jd_text"]), (3.0, 5, "We need Python"))

    def test_apply_job_is_robust_to_odd_data(self):
        state = {}
        app.apply_job_to_state(state, {"jd_text": None, "settings": {"weights": [500, -5], "shortlist": 0, "min_years": None}})
        self.assertEqual((state["w_skill"], state["w_exp"], state["w_kw"]), (100, 0, 0))
        self.assertEqual((state["shortlist"], state["min_years"], state["jd_text"]), (3, 0.0, "None"))
        app.apply_job_to_state(state, {})
        self.assertEqual(state["req_skills"], [])

    def test_audit_summary_and_description(self):
        syn = app.generate_synthetic_candidates(n=200, seed=4, bias_penalty=25)
        sections = app.synthetic_fairness_sections(syn)
        summary = app.audit_summary(sections)
        self.assertEqual(len(summary["sections"]), 2)
        self.assertTrue(all(s["simulated"] for s in summary["sections"]))
        json.dumps(summary)  # must be storable
        self.assertIn("simulated", app.describe_summary("audit", summary))
        self.assertIn("3 resumes", app.describe_summary("screening", {"resumes": 3, "top_score": 90}))
        self.assertEqual(app.describe_summary("audit", {}), "No metrics")

    def test_audit_report_html(self):
        sections = app.synthetic_fairness_sections(app.generate_synthetic_candidates(n=100, seed=1))
        out = app.build_audit_report_html(sections, generated_on="2026-01-01")
        self.assertIn("SIMULATED", out)
        self.assertIn("fairness audit", out)
        self.assertIn("No fairness results", app.build_audit_report_html([]))

    def test_sign_out_clears_everything_except_theme(self):
        state = {"user": {"id": 1}, "dark_mode": True, "intro_done": True, "docs": [1], "jd_text": "secret", "dec_x": "Advance"}
        app.sign_out(state)
        self.assertEqual(state, {"dark_mode": True, "intro_done": True})


class PrototypeAccountTests(unittest.TestCase):
    def test_app_never_talks_to_a_database_directly(self):
        import inspect

        source = inspect.getsource(app)
        for forbidden in ("sqlite3", "psycopg2", "pbkdf2", "smtplib", "hashlib.scrypt", "open("):
            self.assertNotIn(forbidden, source, forbidden)  # storage and crypto live only in auth.py

    def test_prototype_workspace_never_touches_a_database(self):
        state = {}
        ws = app.SessionWorkspace(state)
        self.assertIsNone(ws.add_job("My job", "jd", {"required": ["python"]}))
        self.assertIsNone(ws.add_report("audit", "A", "<p>x</p>", {}))
        self.assertFalse(ws.persistent)
        self.assertEqual(ws.list_jobs()[0]["title"], "My job")
        self.assertEqual(set(state), {"saved_jobs", "saved_reports", "next_id"})

    def test_password_strength_and_meter(self):
        self.assertEqual(app.password_strength("")["label"], "Enter a password")
        self.assertEqual(app.password_strength("short1A!")["label"], "Too short")
        self.assertEqual(app.password_strength("alllowercaseletters")["label"], "Weak")
        self.assertEqual(app.password_strength("Correct-Horse-42")["label"], "Strong")
        self.assertEqual(app.password_strength("Correcthorse42")["label"], "Good")
        for pw in ["", "abc", "Correct-Horse-42", "<script>alert(1)</script>", "$$$ $$$"]:
            out = app.password_meter_html(pw)
            self.assertNotIn("\n", out)
            self.assertNotIn("<script>", out)
            if pw:
                self.assertNotIn(pw, out)  # the meter never echoes the password back
        self.assertIn("width:0%", app.password_meter_html(""))
        self.assertIn("width:100%", app.password_meter_html("Correct-Horse-42"))

    def test_validation_messages(self):
        self.assertEqual(app.prototype_signup_errors("a@b.co", "Correct-Horse-42", "Correct-Horse-42", True), [])
        errors = app.prototype_signup_errors("nope", "short", "different", False)
        self.assertEqual(len(errors), 4)
        self.assertEqual(app.prototype_signin_errors("a@b.co", "x"), [])
        self.assertEqual(len(app.prototype_signin_errors("", "")), 2)
        self.assertEqual(app.prototype_password_errors("Correct-Horse-42", "Correct-Horse-42"), [])
        self.assertTrue(app.simple_email_ok("a.b+c@example.co.uk"))
        self.assertFalse(app.simple_email_ok("a@b"))
        self.assertFalse(app.simple_email_ok(None))

    def test_roles_and_demo_user_have_no_password(self):
        self.assertEqual(app.role_for_email("Admin.demo@example.com"), "admin")
        self.assertEqual(app.role_for_email("asha@example.com"), "recruiter")
        user = app.make_demo_user("  Boss@Example.com ", "admin")
        self.assertEqual((user["email"], user["role"]), ("boss@example.com", "admin"))
        self.assertEqual(app.make_demo_user("a@b.co", "root")["role"], "recruiter")
        self.assertNotIn("password", json.dumps(user).lower())

    def test_start_session_seeds_sample_data_once(self):
        state = {}
        app._start_session(state, "a@example.com", "recruiter")
        self.assertEqual(state["user"]["email"], "a@example.com")
        self.assertEqual(len(state["saved_jobs"]), 2)
        self.assertEqual(sorted(r["kind"] for r in state["saved_reports"]), ["audit", "screening"])
        self.assertTrue(all("(sample)" in x["title"] for x in state["saved_jobs"] + state["saved_reports"]))
        self.assertEqual(len(state["demo_users"]), 3)
        app._start_session(state, "a@example.com", "recruiter")
        self.assertEqual(len(state["saved_jobs"]), 2)  # not seeded twice
        self.assertNotIn("password", json.dumps({k: v for k, v in state.items() if k != "saved_reports"}).lower())

    def test_saved_items_are_session_only_and_limited(self):
        state = {}
        self.assertIsNone(app.add_saved_job(state, "My job", "jd", {"required": ["python"]}))
        self.assertIsNotNone(app.add_saved_job(state, "  ", "jd", {}))
        self.assertIsNotNone(app.add_saved_job(state, "x" * 121, "jd", {}))
        job_id = state["saved_jobs"][0]["id"]
        self.assertTrue(app.delete_saved(state, "saved_jobs", job_id))
        self.assertFalse(app.delete_saved(state, "saved_jobs", job_id))
        self.assertIsNone(app.add_saved_report(state, "audit", "A", "<p>x</p>", {}))
        self.assertIsNotNone(app.add_saved_report(state, "weird", "A", "<p>x</p>", {}))
        self.assertIsNotNone(app.add_saved_report(state, "audit", "A", "x" * (app.MAX_SAVED_REPORT_BYTES + 1), {}))
        for i in range(app.MAX_SAVED_ITEMS):
            app.add_saved_job(state, f"job {i}", "", {})
        self.assertIn("up to", app.add_saved_job(state, "one too many", "", {}))

    def test_admin_actions_only_touch_fake_accounts(self):
        state = {}
        app.seed_demo_data(state)
        email = state["demo_users"][0]["email"]
        self.assertIn("administrator", app.demo_user_action(state, email, "Make administrator"))
        self.assertEqual(state["demo_users"][0]["role"], "admin")
        self.assertIn("recruiter", app.demo_user_action(state, email, "Make recruiter"))
        self.assertIn("Nothing was changed", app.demo_user_action(state, email, "Reset password"))
        self.assertIn("deleted", app.demo_user_action(state, email, "Delete account"))
        self.assertEqual(app.demo_user_action(state, email, "Delete account"), "Account not found.")
        self.assertEqual(app.demo_user_action(state, state["demo_users"][0]["email"], "Explode"), "Unknown action.")
        self.assertTrue(all(u["email"].endswith("@example.com") for u in app.SAMPLE_USERS))


class RealResumeIngestTests(unittest.TestCase):
    def test_helper_pdf_is_readable(self):
        pages = app.read_pdf_pages(make_pdf([R1_P1, R1_P2]))
        self.assertEqual(len(pages), 2)
        self.assertIn("priya.sharma@example.com", pages[0])
        self.assertIn("Tableau", pages[1])

    def test_clean_extracted_text(self):
        raw = "Pro\ufb01cient  in\u00a0Python\u200b \n\n\n\n  Next   line \x00"
        cleaned = app.clean_extracted_text(raw)
        self.assertEqual(cleaned, "Proficient in Python\n\nNext line")
        self.assertEqual(app.clean_extracted_text(None), "")

    def test_auto_split_finds_each_resume(self):
        pages = [R1_P1, R1_P2, R2_P1, R3_P1]
        self.assertEqual(app.detect_resume_starts(pages), [0, 2, 3])
        self.assertEqual(app.split_pdf_pages(pages, "auto"), [(0, 1), (2, 2), (3, 3)])

    def test_running_header_does_not_split_a_resume(self):
        page2 = "priya.sharma@example.com\n" + R1_P2
        self.assertEqual(app.detect_resume_starts([R1_P1, page2, R2_P1]), [0, 2])

    def test_same_phone_on_every_page_stays_together(self):
        page2 = "Priya Sharma | +91 98765 43210\nExperience\nMore work"
        self.assertEqual(app.detect_resume_starts([R1_P1, page2]), [0])

    def test_dates_are_not_mistaken_for_phone_numbers(self):
        self.assertEqual(app._page_identities("Experience\n2018 - 2022 2023 2024"), set())

    def test_split_modes(self):
        pages = [R1_P1, R1_P2, "", R2_P1]
        self.assertEqual(app.split_pdf_pages(pages, "single"), [(0, 3)])
        self.assertEqual(app.split_pdf_pages(pages, "per_page"), [(0, 0), (1, 1), (3, 3)])
        self.assertEqual(app.split_pdf_pages([], "auto"), [])
        self.assertEqual(app.detect_resume_starts([]), [])

    def test_ingest_multi_resume_pdf(self):
        pdf = make_pdf([R1_P1, R1_P2, R2_P1, R3_P1])
        documents, rows, problems = app.ingest_uploads([("batch.pdf", pdf)], "auto")
        self.assertEqual(problems, [])
        self.assertEqual([d[0] for d in documents], ["batch.pdf (pages 1-2)", "batch.pdf (page 3)", "batch.pdf (page 4)"])
        self.assertIn("Tableau", documents[0][1])
        self.assertEqual((rows[0]["pages"], rows[0]["resumes"], rows[0]["detail"]), (4, 3, "1-2, 3, 4"))
        single, rows, _ = app.ingest_uploads([("batch.pdf", pdf)], "single")
        self.assertEqual((len(single), single[0][0]), (1, "batch.pdf"))
        per_page, _, _ = app.ingest_uploads([("batch.pdf", pdf)], "per_page")
        self.assertEqual(len(per_page), 4)
        self.assertEqual(len(app.ingest_uploads([("batch.pdf", pdf)], "nonsense")[0]), 3)  # falls back to auto

    def test_ingest_mixed_files_and_problems(self):
        good_pdf = make_pdf([R2_P1])
        files = [
            ("one.pdf", good_pdf),
            ("notes.txt", b"Python developer with SQL"),
            ("empty.txt", b""),
            ("bad.pdf", b"not a pdf"),
            ("virus.exe", b"abc"),
            ("huge.pdf", b"a" * (app.MAX_FILE_BYTES + 1)),
        ]
        documents, rows, problems = app.ingest_uploads(files, "auto")
        self.assertEqual([d[0] for d in documents], ["one.pdf", "notes.txt"])
        self.assertEqual(len(problems), 4)
        self.assertEqual({r["file"]: r["status"] for r in rows}["bad.pdf"], "Skipped")
        self.assertEqual(sum(1 for r in rows if r["status"] == "Read"), 2)

    def test_scanned_pdf_gives_a_clear_message(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        documents, rows, problems = app.ingest_uploads([("scan.pdf", buffer.getvalue())], "auto")
        self.assertEqual(documents, [])
        self.assertIn("scanned images are not supported", problems[0])
        self.assertEqual(rows[0]["status"], "Skipped")

    def test_page_and_resume_limits(self):
        old = app.MAX_RESUMES_PER_RUN
        app.MAX_RESUMES_PER_RUN = 2
        try:
            docs, _, problems = app.ingest_uploads([("a.txt", b"Python"), ("b.txt", b"SQL"), ("c.txt", b"Java")], "auto")
        finally:
            app.MAX_RESUMES_PER_RUN = old
        self.assertEqual(len(docs), 2)
        self.assertTrue(any("Only the first 2" in p for p in problems))
        old_pages = app.MAX_PDF_PAGES
        app.MAX_PDF_PAGES = 1
        try:
            docs, _, problems = app.ingest_uploads([("long.pdf", make_pdf([R1_P1, R1_P2]))], "auto")
        finally:
            app.MAX_PDF_PAGES = old_pages
        self.assertEqual(docs, [])
        self.assertIn("more than 1 pages", problems[0])

    def test_blank_pages_in_per_page_mode_are_reported(self):
        docs, _, problems = app.ingest_uploads([("x.pdf", make_pdf([R1_P1, " ", R2_P1]))], "per_page")
        self.assertEqual(len(docs), 2)
        self.assertTrue(any("blank page" in p for p in problems))

    def test_end_to_end_screening_of_one_multi_resume_pdf(self):
        documents, _, _ = app.ingest_uploads([("batch.pdf", make_pdf([R1_P1, R1_P2, R2_P1, R3_P1]))], "auto")
        results, warnings = app.screen_resumes(documents, ["python", "sql", "machine learning"], "", 3)
        self.assertEqual(warnings, [])
        self.assertEqual(len(results), 3)
        self.assertEqual(results[0]["file_name"], "batch.pdf (pages 1-2)")  # the Python + SQL + ML resume wins
        self.assertTrue(all("@example" not in r["text_used"] for r in results))
        self.assertTrue(all("[NAME]" in r["text_used"] for r in results))

    def test_ingest_report_html(self):
        rows = [
            {"file": "<b>x</b>.pdf", "kind": "PDF", "pages": 4, "resumes": 3, "detail": "1-2, 3, 4", "status": "Read"},
            {"file": "bad.pdf", "kind": "PDF", "pages": "-", "resumes": 0, "detail": "", "status": "Skipped"},
        ]
        out = app.ingest_report_html(rows)
        self.assertIn("3 resume(s) from 2 file(s)", out)
        self.assertNotIn("<b>x</b>", out)
        self.assertNotIn("\n", out)
        self.assertIn(" open>", out)  # multi-resume file: shown expanded
        self.assertNotIn(" open>", app.ingest_report_html([dict(rows[0], resumes=1)]))
        self.assertEqual(app.ingest_report_html([]), "")


class BlindReviewCardTests(unittest.TestCase):
    def setUp(self):
        text = "Priya Sharma\npriya.sharma@example.com\nPython and SQL, 5 years of experience."
        self.item = app.screen_resumes([("f.txt", text)], ["python"], "", 0)[0][0]

    def test_default_card_never_contains_the_original_identity(self):
        out = app.candidate_card_html(self.item)
        for secret in ("Priya", "Sharma", "priya.sharma@example.com"):
            self.assertNotIn(secret, out)
        self.assertIn("Text used for scoring", out)
        self.assertIn("[NAME]", out)
        self.assertNotIn('class="fh-del"', out)

    def test_original_is_shown_only_on_request(self):
        out = app.candidate_card_html(self.item, show_original=True)
        self.assertIn("Priya", out)
        self.assertIn('class="fh-del"', out)
        self.assertIn("Before / after anonymization", out)

    def test_card_is_safe_and_markdown_proof(self):
        for flag in (False, True):
            out = app.candidate_card_html(self.item, show_original=flag)
            self.assertNotIn("\n", out)
            self.assertNotIn("<script", out)


class RealNameHandlingTests(unittest.TestCase):
    def test_is_name_line(self):
        for good in ["Priya Sharma", "PRIYA SHARMA", "John A. Smith", "Maria de la Cruz", "Anne-Marie O'Neil"]:
            self.assertTrue(app._is_name_line(good), good)
        for bad in ["Python Developer", "Software Engineering", "Work Experience", "Machine Learning", "Skills", "",
                    "Aarav", "jane doe", "Data Scientist at Example Analytics Pvt Ltd"]:
            self.assertFalse(app._is_name_line(bad), bad)

    def test_name_is_found_below_a_title_line_and_masked_everywhere(self):
        text = "Curriculum Vitae\nPriya Sharma\npriya@example.com\nPriya has 5 years of experience. Ms. Sharma led a team. Priyanka is a friend."
        out = app.anonymize_text(text)
        self.assertNotIn("Priya Sharma", out)
        self.assertNotIn("Sharma", out)
        self.assertNotIn("Priya ", out)
        self.assertIn("Priyanka", out)  # a different word is left alone
        self.assertTrue(out.startswith("Curriculum Vitae"))

    def test_name_label_value_is_masked_in_the_body_too(self):
        out = app.anonymize_text("Name: Rahul Verma\nBackend developer. Rahul built APIs.")
        self.assertNotIn("Rahul", out)
        self.assertNotIn("Verma", out)

    def test_name_parts_that_are_skills_or_titles_are_not_masked(self):
        tokens = app._name_tokens("Python Manager Smith")
        self.assertIn("Smith", tokens)
        self.assertNotIn("Python", tokens)
        self.assertNotIn("Manager", tokens)

    def test_detect_name_scans_the_first_lines(self):
        self.assertEqual(app.detect_name("Resume\nPriya Sharma\nSkills"), "Priya Sharma")
        self.assertIsNone(app.detect_name("Skills\nPython and SQL"))


class FakePostgresServer:
    """A stand-in for a Postgres server, built on SQLite, so the Postgres code path can be exercised.

    It insists on Postgres-style SQL (``%s`` placeholders, no AUTOINCREMENT, RETURNING for inserts and
    no use of ``lastrowid``). It is NOT a real Postgres server.
    """

    class OperationalError(Exception):
        pass

    class InterfaceError(Exception):
        pass

    class IntegrityError(Exception):
        pass

    def __init__(self):
        self.store = sqlite3.connect(":memory:", check_same_thread=False)
        self.store.row_factory = sqlite3.Row
        self.store.execute("PRAGMA foreign_keys = ON")
        self.statements = []
        self.fail_next = 0
        self.connects = 0
        server = self

        class Cursor:
            def __init__(self, conn):
                self.conn = conn
                self._rows = []
                self._rowcount = -1

            def execute(self, sql, params=()):
                if server.fail_next:
                    server.fail_next -= 1
                    raise server.OperationalError("connection lost")
                assert "?" not in sql, "placeholders must be translated to %s: " + sql
                assert "AUTOINCREMENT" not in sql and "REAL NOT NULL" not in sql.replace("DOUBLE PRECISION", ""), sql
                server.statements.append(sql)
                text = (
                    sql.replace("%s", "?").replace("SERIAL PRIMARY KEY", "INTEGER PRIMARY KEY AUTOINCREMENT")
                    .replace("DOUBLE PRECISION", "REAL")
                )
                try:
                    cursor = server.store.execute(text, tuple(params))
                    self._rows = cursor.fetchall()  # read RETURNING rows before committing (SQLite requires it)
                    self._rowcount = cursor.rowcount
                    server.store.commit()
                except sqlite3.IntegrityError as exc:
                    server.store.rollback()
                    raise server.IntegrityError(str(exc)) from exc

            def fetchall(self):
                return list(self._rows)

            def fetchone(self):
                return self._rows[0] if self._rows else None

            @property
            def rowcount(self):
                return self._rowcount

            def close(self):
                pass

        class Connection:
            autocommit = False
            closed = False

            def cursor(self, cursor_factory=None):
                return Cursor(self)

            def close(self):
                self.closed = True

        class Extras:
            RealDictCursor = object()

        class Module:
            OperationalError = server.OperationalError
            InterfaceError = server.InterfaceError
            IntegrityError = server.IntegrityError

            @staticmethod
            def connect(dsn, connect_timeout=None):
                server.connects += 1
                return Connection()

        self.module = Module
        self.extras = Extras


class AuthTestCase(unittest.TestCase):
    """Fresh database per test and cheap hashing. ``backend`` selects SQLite or the fake Postgres."""

    backend = "sqlite"

    def setUp(self):
        self._old_iterations = auth.HASH_ITERATIONS
        auth.HASH_ITERATIONS = 1000
        auth._DUMMY_HASH = None
        self._patch = None
        if self.backend == "postgres":
            from unittest import mock

            self.server = FakePostgresServer()
            self._patch = mock.patch.object(auth, "_import_psycopg2", lambda: (self.server.module, self.server.extras))
            self._patch.start()
            self.conn = auth.Database("postgres", "postgresql://user:pw@host/db")
        else:
            self.conn = auth.Database("sqlite", ":memory:")
        self.pw = "Correct-Horse-42"

    def tearDown(self):
        if self._patch:
            self._patch.stop()
        auth.HASH_ITERATIONS = self._old_iterations
        auth._DUMMY_HASH = None
        self.conn.close()

    def make_user(self, email, **kwargs):
        user, error = auth.create_user(self.conn, email, kwargs.pop("password", self.pw), **kwargs)
        self.assertIsNone(error, error)
        return user


class PasswordAndValidationTests(unittest.TestCase):
    def setUp(self):
        self._old = auth.HASH_ITERATIONS
        auth.HASH_ITERATIONS = 1000

    def tearDown(self):
        auth.HASH_ITERATIONS = self._old

    def test_password_policy(self):
        self.assertIsNotNone(auth.validate_password("short1A"))
        self.assertIsNotNone(auth.validate_password("1234567890"))
        self.assertIsNotNone(auth.validate_password("password123"))
        self.assertIsNotNone(auth.validate_password("alllowercaseletters"))
        self.assertIsNotNone(auth.validate_password("x" * 129))
        self.assertIsNotNone(auth.validate_password("Jonathan-2024!", "jonathan@example.com"))
        self.assertIsNone(auth.validate_password("Correct-Horse-42", "jo@example.com"))

    def test_email_validation_and_normalisation(self):
        self.assertIsNone(auth.validate_email("a.b+c@example.co.uk"))
        for bad in ["", "no-at-sign", "a@b", "a@@b.com", "a b@c.com", "x" * 250 + "@a.com"]:
            self.assertIsNotNone(auth.validate_email(bad), bad)
        self.assertEqual(auth.normalize_email("  Bob@Example.COM "), "bob@example.com")

    def test_hash_and_verify(self):
        stored = auth.hash_password("Correct-Horse-42")
        self.assertTrue(stored.startswith("pbkdf2_sha256$"))
        self.assertNotIn("Correct-Horse-42", stored)
        self.assertTrue(auth.verify_password("Correct-Horse-42", stored))
        self.assertFalse(auth.verify_password("wrong", stored))
        self.assertNotEqual(stored, auth.hash_password("Correct-Horse-42"))  # random salt
        for broken in ["", "garbage", "a$b$c$d", "pbkdf2_sha256$x$zz$yy"]:
            self.assertFalse(auth.verify_password("x", broken))


class UserAndLoginTests(AuthTestCase):
    def test_first_user_is_admin_then_recruiters(self):
        first = self.make_user("Boss@Example.com")
        second = self.make_user("rec@example.com")
        self.assertEqual(first["role"], "admin")
        self.assertEqual(first["email"], "boss@example.com")
        self.assertEqual(second["role"], "recruiter")

    def test_admin_email_rule(self):
        a = self.make_user("rec@example.com", admin_email="boss@example.com")
        b = self.make_user("boss@example.com", admin_email="boss@example.com")
        self.assertEqual((a["role"], b["role"]), ("recruiter", "admin"))

    def test_duplicate_and_weak_signups_rejected(self):
        self.make_user("a@example.com")
        self.assertIn("already exists", auth.create_user(self.conn, "A@example.com", self.pw)[1])
        self.assertIsNotNone(auth.create_user(self.conn, "b@example.com", "weak")[1])
        self.assertIsNotNone(auth.create_user(self.conn, "not-an-email", self.pw)[1])
        self.assertIsNotNone(auth.create_user(self.conn, "c@example.com", self.pw, role="root")[1])

    def test_passwords_are_not_stored_in_plain_text(self):
        self.make_user("a@example.com")
        row = self.conn.one("SELECT password_hash FROM users")
        self.assertNotIn(self.pw, row["password_hash"])

    def test_login_success_and_generic_errors(self):
        self.make_user("a@example.com")
        user, error = auth.authenticate(self.conn, " A@Example.com ", self.pw)
        self.assertIsNone(error)
        self.assertEqual(user["email"], "a@example.com")
        _, wrong = auth.authenticate(self.conn, "a@example.com", "nope")
        _, unknown = auth.authenticate(self.conn, "nobody@example.com", "nope")
        self.assertEqual(wrong, unknown)  # cannot tell which emails exist
        self.assertEqual(wrong, auth.GENERIC_LOGIN_ERROR)

    def test_lockout_and_expiry(self):
        self.make_user("a@example.com")
        now = 1_700_000_000.123
        for _ in range(auth.MAX_FAILED_ATTEMPTS):
            auth.authenticate(self.conn, "a@example.com", "bad", now=now)
        user, error = auth.authenticate(self.conn, "a@example.com", self.pw, now=now + 10)
        self.assertIsNone(user)
        self.assertIn("Too many", error)
        user, error = auth.authenticate(self.conn, "a@example.com", self.pw, now=now + auth.LOCK_SECONDS + 1)
        self.assertIsNone(error)

    def test_unverified_user_cannot_sign_in(self):
        self.make_user("a@example.com", verified=False)
        user, error = auth.authenticate(self.conn, "a@example.com", self.pw)
        self.assertIsNone(user)
        self.assertIn("verify", error)
        auth.set_verified(self.conn, auth.get_user_by_email(self.conn, "a@example.com")["id"])
        self.assertIsNotNone(auth.authenticate(self.conn, "a@example.com", self.pw)[0])

    def test_sql_injection_attempts_are_harmless(self):
        self.make_user("a@example.com")
        user, error = auth.authenticate(self.conn, "' OR '1'='1", "' OR '1'='1")
        self.assertIsNone(user)
        user, _ = auth.authenticate(self.conn, "a@example.com' --", self.pw)
        self.assertIsNone(user)
        job_id, _ = auth.save_job(self.conn, 1, "x'); DROP TABLE users; --", "jd", {})
        self.assertIsNotNone(job_id)
        self.assertEqual(auth.count_users(self.conn), 1)


class CodeAndPasswordChangeTests(AuthTestCase):
    def test_code_flow(self):
        u = self.make_user("a@example.com")
        code = auth.issue_code(self.conn, u["id"], "reset", now=100.0)
        self.assertRegex(code, r"^\d{6}$")
        self.assertFalse(auth.check_code(self.conn, u["id"], "verify", code, now=101.0))  # wrong purpose
        self.assertTrue(auth.check_code(self.conn, u["id"], "reset", code, now=101.0))
        self.assertFalse(auth.check_code(self.conn, u["id"], "reset", code, now=102.0))  # single use

    def test_code_expiry_and_attempt_limit(self):
        u = self.make_user("a@example.com")
        code = auth.issue_code(self.conn, u["id"], "reset", now=100.0)
        self.assertFalse(auth.check_code(self.conn, u["id"], "reset", code, now=100.0 + auth.CODE_TTL_SECONDS + 1))
        code = auth.issue_code(self.conn, u["id"], "reset", now=100.0)
        wrong = "000000" if code != "000000" else "111111"
        for _ in range(auth.MAX_CODE_ATTEMPTS):
            self.assertFalse(auth.check_code(self.conn, u["id"], "reset", wrong, now=101.0))
        self.assertFalse(auth.check_code(self.conn, u["id"], "reset", code, now=102.0))  # locked out

    def test_change_password(self):
        u = self.make_user("a@example.com")
        self.assertIn("not correct", auth.change_password(self.conn, u["id"], "wrong", "Another-Pass-77"))
        self.assertIn("different", auth.change_password(self.conn, u["id"], self.pw, self.pw))
        self.assertIsNotNone(auth.change_password(self.conn, u["id"], self.pw, "weak"))
        self.assertIsNone(auth.change_password(self.conn, u["id"], self.pw, "Another-Pass-77"))
        self.assertIsNone(auth.authenticate(self.conn, "a@example.com", "Another-Pass-77")[1])
        self.assertIsNotNone(auth.authenticate(self.conn, "a@example.com", self.pw)[1])

    def test_set_password_validates_and_clears_flags(self):
        u = self.make_user("a@example.com", must_change=True)
        self.assertTrue(u["must_change"])
        self.assertIsNotNone(auth.set_password(self.conn, u["id"], "short"))
        self.assertIsNone(auth.set_password(self.conn, u["id"], "Brand-New-Pass-9"))
        self.assertFalse(auth.get_user(self.conn, u["id"])["must_change"])
        self.assertEqual(auth.set_password(self.conn, 999, "Brand-New-Pass-9"), "Account not found.")


class RoleTests(AuthTestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.make_user("admin@example.com")
        self.rec = self.make_user("rec@example.com")
        self.other = self.make_user("other@example.com")

    def test_recruiter_cannot_use_admin_functions(self):
        for call in (
            lambda: auth.list_users(self.conn, self.rec["id"]),
            lambda: auth.set_role(self.conn, self.rec["id"], self.other["id"], "admin"),
            lambda: auth.delete_user(self.conn, self.rec["id"], self.other["id"]),
            lambda: auth.admin_reset_password(self.conn, self.rec["id"], self.other["id"]),
        ):
            with self.assertRaises(PermissionError):
                call()
        with self.assertRaises(PermissionError):
            auth.list_users(self.conn, 9999)  # unknown actor

    def test_role_change_takes_effect_immediately(self):
        self.assertIsNone(auth.set_role(self.conn, self.admin["id"], self.rec["id"], "admin"))
        self.assertEqual(len(auth.list_users(self.conn, self.rec["id"])), 3)
        self.assertIsNone(auth.set_role(self.conn, self.admin["id"], self.rec["id"], "recruiter"))
        with self.assertRaises(PermissionError):
            auth.list_users(self.conn, self.rec["id"])

    def test_admin_guard_rails(self):
        self.assertIn("own role", auth.set_role(self.conn, self.admin["id"], self.admin["id"], "recruiter"))
        self.assertIn("own account", auth.delete_user(self.conn, self.admin["id"], self.admin["id"]))
        self.assertIsNotNone(auth.set_role(self.conn, self.admin["id"], self.rec["id"], "root"))
        self.assertIsNotNone(auth.set_role(self.conn, self.admin["id"], 999, "admin"))

    def test_admin_sees_counts_not_content(self):
        auth.save_job(self.conn, self.rec["id"], "Secret job", "secret description", {})
        auth.save_report(self.conn, self.rec["id"], "Secret report", "screening", "<p>secret</p>", {})
        rows = {u["email"]: u for u in auth.list_users(self.conn, self.admin["id"])}
        self.assertEqual((rows["rec@example.com"]["jobs"], rows["rec@example.com"]["reports"]), (1, 1))
        for row in rows.values():
            self.assertNotIn("secret", json.dumps(row).lower())
        self.assertNotIn("password_hash", rows["rec@example.com"])
        report_id = auth.list_reports(self.conn, self.rec["id"])[0]["id"]
        self.assertIsNone(auth.get_report_html(self.conn, self.admin["id"], report_id))

    def test_admin_reset_password_forces_change(self):
        temporary, error = auth.admin_reset_password(self.conn, self.admin["id"], self.rec["id"])
        self.assertIsNone(error)
        user, err = auth.authenticate(self.conn, "rec@example.com", temporary)
        self.assertIsNone(err)
        self.assertTrue(user["must_change"])
        self.assertIsNotNone(auth.authenticate(self.conn, "rec@example.com", self.pw)[1])
        self.assertEqual(auth.admin_reset_password(self.conn, self.admin["id"], 999)[1], "Account not found.")

    def test_delete_user_removes_their_data(self):
        auth.save_job(self.conn, self.rec["id"], "J", "jd", {})
        auth.save_report(self.conn, self.rec["id"], "R", "audit", "<p>x</p>", {})
        self.assertIsNone(auth.delete_user(self.conn, self.admin["id"], self.rec["id"]))
        self.assertIsNone(auth.get_user(self.conn, self.rec["id"]))
        self.assertEqual(self.conn.one("SELECT COUNT(*) AS n FROM jobs")["n"], 0)
        self.assertEqual(self.conn.one("SELECT COUNT(*) AS n FROM reports")["n"], 0)


class UserDataIsolationTests(AuthTestCase):
    def setUp(self):
        super().setUp()
        self.a = self.make_user("a@example.com")
        self.b = self.make_user("b@example.com")

    def test_jobs_are_private(self):
        job_id, error = auth.save_job(self.conn, self.a["id"], "Data Scientist", "jd text", {"required": ["python"], "min_years": 2})
        self.assertIsNone(error)
        job = auth.get_job(self.conn, self.a["id"], job_id)
        self.assertEqual(job["settings"]["required"], ["python"])
        self.assertIsNone(auth.get_job(self.conn, self.b["id"], job_id))
        self.assertFalse(auth.delete_job(self.conn, self.b["id"], job_id))
        self.assertEqual(auth.list_jobs(self.conn, self.b["id"]), [])
        self.assertTrue(auth.delete_job(self.conn, self.a["id"], job_id))

    def test_reports_are_private(self):
        report_id, error = auth.save_report(self.conn, self.a["id"], "Q3 screening", "screening", "<p>hi</p>", {"n": 3})
        self.assertIsNone(error)
        self.assertEqual(auth.get_report_html(self.conn, self.a["id"], report_id), "<p>hi</p>")
        self.assertIsNone(auth.get_report_html(self.conn, self.b["id"], report_id))
        self.assertFalse(auth.delete_report(self.conn, self.b["id"], report_id))
        self.assertEqual(auth.list_reports(self.conn, self.b["id"]), [])
        listed = auth.list_reports(self.conn, self.a["id"], kind="screening")
        self.assertEqual(listed[0]["summary"], {"n": 3})
        self.assertEqual(auth.list_reports(self.conn, self.a["id"], kind="audit"), [])
        self.assertTrue(auth.delete_report(self.conn, self.a["id"], report_id))

    def test_validation_and_limits(self):
        self.assertIsNotNone(auth.save_job(self.conn, self.a["id"], "  ", "jd", {})[1])
        self.assertIsNotNone(auth.save_job(self.conn, self.a["id"], "x" * 121, "jd", {})[1])
        self.assertIsNotNone(auth.save_job(self.conn, self.a["id"], "t", "x" * (auth.MAX_JD_CHARS + 1), {})[1])
        self.assertIsNotNone(auth.save_report(self.conn, self.a["id"], "t", "weird", "<p/>", {})[1])
        self.assertIsNotNone(auth.save_report(self.conn, self.a["id"], "t", "audit", "x" * (auth.MAX_REPORT_BYTES + 1), {})[1])
        old = auth.MAX_JOBS_PER_USER
        auth.MAX_JOBS_PER_USER = 1
        try:
            self.assertIsNone(auth.save_job(self.conn, self.a["id"], "one", "jd", {})[1])
            self.assertIn("up to", auth.save_job(self.conn, self.a["id"], "two", "jd", {})[1])
            self.assertIsNone(auth.save_job(self.conn, self.b["id"], "mine", "jd", {})[1])  # per-user cap
        finally:
            auth.MAX_JOBS_PER_USER = old


class UserAndLoginTestsPostgres(UserAndLoginTests):
    backend = "postgres"


class CodeAndPasswordChangeTestsPostgres(CodeAndPasswordChangeTests):
    backend = "postgres"


class RoleTestsPostgres(RoleTests):
    backend = "postgres"


class UserDataIsolationTestsPostgres(UserDataIsolationTests):
    backend = "postgres"


class DatabaseLayerTests(unittest.TestCase):
    def setUp(self):
        self._old = auth.HASH_ITERATIONS
        auth.HASH_ITERATIONS = 1000

    def tearDown(self):
        auth.HASH_ITERATIONS = self._old

    def test_data_survives_a_restart_sqlite_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "fairhire.db")
            first = auth.Database("sqlite", path)
            user, _ = auth.create_user(first, "a@example.com", "Correct-Horse-42")
            auth.save_job(first, user["id"], "Kept job", "jd", {"required": ["python"]})
            first.close()
            second = auth.Database("sqlite", path)  # a "restarted" app opening the same file
            self.assertIsNone(auth.authenticate(second, "a@example.com", "Correct-Horse-42")[1])
            self.assertEqual(auth.list_jobs(second, user["id"])[0]["title"], "Kept job")
            second.close()

    def test_data_survives_a_restart_postgres_path(self):
        from unittest import mock

        server = FakePostgresServer()
        with mock.patch.object(auth, "_import_psycopg2", lambda: (server.module, server.extras)):
            first = auth.Database("postgres", "postgresql://u:p@h/db")
            user, _ = auth.create_user(first, "a@example.com", "Correct-Horse-42")
            auth.save_report(first, user["id"], "Kept report", "audit", "<p>x</p>", {"n": 1})
            second = auth.Database("postgres", "postgresql://u:p@h/db")  # a new process, same server
        self.assertIsNone(auth.authenticate(second, "a@example.com", "Correct-Horse-42")[1])
        self.assertEqual(auth.list_reports(second, user["id"])[0]["title"], "Kept report")

    def test_postgres_sql_translation_and_schema(self):
        from unittest import mock

        server = FakePostgresServer()
        with mock.patch.object(auth, "_import_psycopg2", lambda: (server.module, server.extras)):
            db = auth.Database("postgres", "postgresql://u:p@h/db")
            auth.create_user(db, "a@example.com", "Correct-Horse-42")
        joined = "\n".join(server.statements)
        self.assertIn("SERIAL PRIMARY KEY", joined)
        self.assertIn("DOUBLE PRECISION", joined)
        self.assertIn("RETURNING id", joined)
        self.assertNotIn("?", joined)
        self.assertNotIn("AUTOINCREMENT", joined)

    def test_dropped_connection_is_retried_once(self):
        from unittest import mock

        server = FakePostgresServer()
        with mock.patch.object(auth, "_import_psycopg2", lambda: (server.module, server.extras)):
            db = auth.Database("postgres", "postgresql://u:p@h/db")
            auth.create_user(db, "a@example.com", "Correct-Horse-42")
            before = server.connects
            server.fail_next = 1
            self.assertEqual(auth.get_user_by_email(db, "a@example.com")["email"], "a@example.com")
            self.assertEqual(server.connects, before + 1)  # reconnected
            server.fail_next = 2  # fails twice: the error is raised, not hidden
            with self.assertRaises(FakePostgresServer.OperationalError):
                auth.get_user_by_email(db, "a@example.com")

    def test_integrity_errors_use_one_type(self):
        for kind in ("sqlite", "postgres"):
            from unittest import mock

            server = FakePostgresServer()
            with mock.patch.object(auth, "_import_psycopg2", lambda: (server.module, server.extras)):
                db = auth.Database(kind, ":memory:" if kind == "sqlite" else "postgresql://u:p@h/db")
                db.insert("INSERT INTO users (email, password_hash, created_at) VALUES (?, ?, ?)", ("x@example.com", "h", "now"))
                with self.assertRaises(auth.DbIntegrityError):
                    db.insert("INSERT INTO users (email, password_hash, created_at) VALUES (?, ?, ?)", ("x@example.com", "h", "now"))
                self.assertEqual(db.one("SELECT COUNT(*) AS n FROM users")["n"], 1)  # still usable after the error

    def test_factory_caches_files_but_not_memory(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "a.db")
            self.assertIs(auth.get_database(None, path), auth.get_database(None, path))
            auth._DATABASES.clear()
        self.assertIsNot(auth.get_database(None, ":memory:"), auth.get_database(None, ":memory:"))
        with self.assertRaises(ValueError):
            auth.Database("mysql", "x")

    def test_missing_postgres_driver_gives_an_error(self):
        from unittest import mock

        def boom():
            raise ImportError("no psycopg2")

        with mock.patch.object(auth, "_import_psycopg2", boom):
            with self.assertRaises(ImportError):
                auth.get_database("postgresql://u:p@h/db")


class EmailTests(unittest.TestCase):
    CONF = {"host": "smtp.example.com", "port": 465, "username": "u", "password": "p", "sender": "noreply@example.com"}

    def test_configuration_detection(self):
        self.assertTrue(auth.smtp_configured(self.CONF))
        self.assertFalse(auth.smtp_configured({"host": "x"}))
        self.assertFalse(auth.smtp_configured(None))
        self.assertEqual(auth.send_email(None, "a@b.co", "s", "b"), (False, "Email is not configured."))

    def test_send_success_and_failure_do_not_leak_details(self):
        from unittest import mock

        sent = []

        class FakeSMTP:
            def __init__(self, *args, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def login(self, user, password):
                sent.append(("login", user))

            def send_message(self, message):
                sent.append(("send", message["To"], message["Subject"]))

        with mock.patch.object(auth.smtplib, "SMTP_SSL", FakeSMTP):
            self.assertEqual(auth.send_email(self.CONF, "a@b.co", "Hello", "Body"), (True, None))
        self.assertIn(("send", "a@b.co", "Hello"), sent)

        class BrokenSMTP(FakeSMTP):
            def login(self, user, password):
                raise OSError("secret server detail p@ssw0rd")

        with mock.patch.object(auth.smtplib, "SMTP_SSL", BrokenSMTP):
            ok, error = auth.send_email(self.CONF, "a@b.co", "Hello", "Body")
        self.assertFalse(ok)
        self.assertNotIn("secret", error)
        self.assertIn("OSError", error)


class JobDescriptionCheckTests(unittest.TestCase):
    BIASED = (
        "We need a young, energetic rockstar ninja developer. Must be a native English speaker, "
        "well-groomed and able-bodied. He will report to the chairman. Citizens only. Culture fit is key."
    )

    def terms(self, text):
        return {f["term"].lower() for f in app.check_job_description(text)["findings"]}

    def test_flags_each_category(self):
        result = app.check_job_description(self.BIASED)
        cats = {f["category"] for f in result["findings"]}
        for expected in ["Age", "Gender-coded", "Nationality / language", "Appearance", "Disability / health", "Culture"]:
            self.assertIn(expected, cats)
        self.assertEqual(result["risk"], "high")
        self.assertGreaterEqual(result["counts"]["high"], 3)
        starts = [f["start"] for f in result["findings"]]
        self.assertEqual(starts, sorted(starts))

    def test_clean_text_has_no_flags(self):
        clean = "We are hiring a data analyst. You will build reports with SQL and explain results to colleagues."
        result = app.check_job_description(clean)
        self.assertEqual(result["findings"], [])
        self.assertEqual(result["risk"], "none")
        self.assertEqual(app.check_job_description("")["risk"], "none")
        self.assertEqual(app.check_job_description(None)["findings"], [])

    def test_no_false_positives_on_common_words(self):
        text = "Over 10 years of experience helps. Senior engineers mentor others. Competitive salary. Younger Street office."
        terms = self.terms(text)
        self.assertNotIn("over 10 years", " ".join(terms))
        self.assertNotIn("competitive", terms)
        self.assertEqual({"younger"} & terms, {"younger"})  # the word itself is flagged; sentences with 'Over 10' are not
        self.assertEqual(self.terms("The team works well together and the shell is fine."), set())

    def test_age_ceiling_and_age_statements(self):
        self.assertIn("no more than 5 years experience", self.terms("Applicants with no more than 5 years experience."))
        self.assertTrue(self.terms("Candidates aged 25-35 preferred."))
        self.assertTrue(self.terms("Must be under 30 years old."))

    def test_requirements_overload(self):
        text = "\n".join(f"- requirement {i}" for i in range(14))
        cats = {f["category"] for f in app.check_job_description(text)["findings"]}
        self.assertIn("Requirements overload", cats)
        self.assertNotIn("Requirements overload", {f["category"] for f in app.check_job_description("\n".join(f"- r{i}" for i in range(5)))["findings"]})

    def test_gender_coding_labels(self):
        masc = "Aggressive, fearless, dominant and competitive. Independent, driven, ambitious, assertive."
        fem = "Collaborative, supportive, caring, empathetic. We trust and understand each other, warm and polite."
        self.assertEqual(app.gender_coding(masc)["label"], "Leans masculine-coded")
        self.assertEqual(app.gender_coding(fem)["label"], "Leans feminine-coded")
        self.assertEqual(app.gender_coding("Build reports.")["label"], "Neutral")
        self.assertEqual(app.gender_coding("Ambitious and collaborative.")["label"], "Balanced")
        self.assertIn("aggressive", app.gender_coding(masc)["masculine_words"])

    def test_rewrite_applies_only_safe_replacements_and_keeps_case(self):
        text = "Rockstar developer. We need a Ninja and a guru. ASSERTIVE? No: AGGRESSIVE. Hire a young person. He leads."
        result = app.check_job_description(text)
        revised, changes = app.rewrite_job_description(text, result["findings"])
        self.assertGreaterEqual(changes, 4)
        self.assertIn("Skilled professional developer", revised)
        self.assertIn("Expert", revised)
        self.assertIn("specialist", revised)
        self.assertIn("PROACTIVE", revised)
        self.assertIn("young", revised)  # no automatic replacement for age wording
        self.assertIn("He leads", revised)  # pronouns are left for a human
        self.assertEqual(app.rewrite_job_description("", [])[1], 0)

    def test_rewrite_ignores_stale_findings(self):
        findings = app.check_job_description("We need a rockstar.")["findings"]
        revised, changes = app.rewrite_job_description("Completely different text here.", findings)
        self.assertEqual((revised, changes), ("Completely different text here.", 0))

    def test_html_is_escaped_and_markdown_proof(self):
        result = app.check_job_description("<script>alert(1)</script> We need a rockstar & ninja. $5")
        for out in (app.jd_summary_html(result), app.jd_findings_html(result), app.gender_bar_html(result["gender"])):
            self.assertNotIn("<script>", out)
            self.assertNotIn("\n", out)
        self.assertIn("No flagged wording", app.jd_findings_html(app.check_job_description("Hello.")))

    def test_report_mentions_flag_count(self):
        out = app.build_report_html([], {"jd_flags": 7})
        self.assertIn("Job-description wording flags", out)
        self.assertNotIn("Job-description wording flags", app.build_report_html([], {}))


class SkillEvidenceTests(unittest.TestCase):
    TEXT = (
        "Summary\nData scientist with 5 years of experience. Built machine learning models in Python for churn prediction.\n"
        "Wrote SQL queries daily. Also used Python for dashboards and automation scripts.\n"
        "2019 - 2024 Data Scientist at Example Analytics\nEducation\n2014 - 2018 B.Sc University"
    )

    def test_finds_snippets_per_skill_in_different_sentences(self):
        ev = app.find_skill_evidence(self.TEXT, ["python", "sql", "java"])
        self.assertEqual(set(ev), {"python", "sql"})  # java is not in the text
        self.assertEqual(len(ev["python"]), 2)
        self.assertEqual({s["hit"].lower() for s in ev["python"]}, {"python"})
        self.assertIn("churn prediction", ev["python"][0]["after"])
        self.assertEqual(len(app.find_skill_evidence(self.TEXT, ["python"], per_skill=1)["python"]), 1)

    def test_aliases_and_custom_banks(self):
        ev = app.find_skill_evidence("I used scikit-learn and sklearn.", ["scikit-learn"])
        self.assertEqual(len(ev["scikit-learn"]), 1)  # one sentence -> one snippet
        bank = app.merged_skill_bank({"rag": ["rag", "retrieval augmented generation"]})
        ev = app.find_skill_evidence("Built Retrieval Augmented Generation apps.", ["rag"], bank)
        self.assertEqual(ev["rag"][0]["hit"], "Retrieval Augmented Generation")

    def test_long_sentences_are_clipped(self):
        text = "word " * 100 + "Python" + " word" * 100 + "."
        snippet = app.find_skill_evidence(text, ["python"])["python"][0]
        self.assertTrue(snippet["left_cut"] and snippet["right_cut"])
        self.assertLessEqual(len(snippet["before"]), 112)
        self.assertLessEqual(len(snippet["after"]), 112)

    def test_experience_evidence(self):
        ev = app.experience_evidence(self.TEXT, current_year=2026)
        hits = [e["hit"] for e in ev]
        self.assertTrue(any("5 years" in h for h in hits))
        self.assertTrue(any("2019 - 2024" in h for h in hits))
        self.assertFalse(any("2014 - 2018" in h for h in hits))  # education dates are not evidence of experience
        self.assertEqual(app.experience_evidence("", None), [])
        self.assertLessEqual(len(app.experience_evidence(self.TEXT, 2026, limit=1)), 1)

    def test_highlighting_is_escaped_and_merges_overlaps(self):
        out = app.highlight_skills_html("Python & <b>SQL</b>. Python again $5", ["python", "sql"])
        self.assertEqual(out.count('<mark class="fh-hit">'), 3)
        self.assertNotIn("<b>", out)
        self.assertNotIn("$", out)
        self.assertNotIn("\n", app.highlight_skills_html("a\n\nb python", ["python"]))
        self.assertEqual(app.highlight_skills_html("", ["python"]), "")
        overlapping = app.highlight_skills_html("machine learning", ["machine learning", "ml"], None)
        self.assertEqual(overlapping.count("<mark"), 1)

    def test_screening_attaches_evidence_from_the_anonymized_text(self):
        text = "Priya Sharma\npriya@example.com\nBuilt models in Python. Priya also wrote SQL. 4 years of experience."
        item = app.screen_resumes([("a.txt", text)], ["python", "sql", "java"], "", 3)[0][0]
        self.assertEqual(set(item["evidence"]), {"python", "sql"})
        blob = json.dumps(item["evidence"]) + json.dumps(item["experience_evidence"])
        self.assertNotIn("Priya", blob)  # evidence comes from the anonymized text
        self.assertTrue(item["experience_evidence"])

    def test_card_shows_evidence_highlights_and_missing_skills(self):
        text = "Jane Doe\nBuilt models in Python. Wrote SQL daily. 4 years of experience."
        item = app.screen_resumes([("a.txt", text)], ["python", "sql", "java"], "", 3)[0][0]
        out = app.candidate_card_html(item)
        self.assertIn("Evidence from the resume", out)
        self.assertIn('class="fh-hit"', out)
        self.assertIn("Not found anywhere in the resume text: java", out)
        self.assertNotIn("Jane", out)
        self.assertNotIn("\n", out)
        self.assertIn("No evidence to show", app.evidence_html({}))

    def test_saved_reports_do_not_contain_resume_snippets(self):
        text = "Built models in Python for churn prediction at Zorblax Corp."
        results, _ = app.screen_resumes([("a.txt", text)], ["python"], "", 0)
        html_out = app.build_report_html(results, {"required": ["python"], "shortlist": 1})
        self.assertNotIn("churn prediction", html_out)
        self.assertNotIn("Zorblax", html_out)


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self._old = auth.HASH_ITERATIONS
        auth.HASH_ITERATIONS = 1000
        self.db = auth.Database("sqlite", ":memory:")
        self.alice, _ = auth.create_user(self.db, "alice@example.com", "Correct-Horse-42")
        self.bob, _ = auth.create_user(self.db, "bob@example.com", "Correct-Horse-42")

    def tearDown(self):
        auth.HASH_ITERATIONS = self._old
        self.db.close()

    def check_workspace(self, ws):
        self.assertIsNone(ws.add_job("Data Scientist", "We need Python", {"required": ["python"], "min_years": 2}))
        self.assertIsNotNone(ws.add_job("  ", "jd", {}))
        job = ws.list_jobs()[0]
        self.assertEqual(job["title"], "Data Scientist")
        full = ws.get_job(job["id"])
        self.assertEqual((full["jd_text"], full["settings"]["required"]), ("We need Python", ["python"]))
        self.assertIsNone(ws.add_report("screening", "Oct", "<p>secret html</p>", {"resumes": 3, "top_score": 90.0}))
        self.assertIsNone(ws.add_report("audit", "Audit", "<p>a</p>", {"sections": []}))
        self.assertIsNotNone(ws.add_report("weird", "x", "<p/>", {}))
        self.assertEqual([r["kind"] for r in ws.list_reports()], ["audit", "screening"])
        self.assertEqual([r["title"] for r in ws.list_reports("screening")], ["Oct"])
        self.assertTrue(all("html" not in r for r in ws.list_reports()))
        rid = ws.list_reports("screening")[0]["id"]
        self.assertEqual(ws.get_report_html(rid), "<p>secret html</p>")
        self.assertIsNone(ws.get_report_html(99999))
        self.assertTrue(ws.delete_report(rid))
        self.assertFalse(ws.delete_report(rid))
        self.assertTrue(ws.delete_job(job["id"]))
        self.assertEqual(ws.list_jobs(), [])

    def test_session_workspace(self):
        self.check_workspace(app.SessionWorkspace({}))

    def test_db_workspace(self):
        self.check_workspace(app.DbWorkspace(self.db, self.alice["id"], {}))

    def test_db_workspace_is_private_and_persistent(self):
        state_a, state_b = {}, {}
        a, b = app.DbWorkspace(self.db, self.alice["id"], state_a), app.DbWorkspace(self.db, self.bob["id"], state_b)
        self.assertTrue(a.persistent)
        a.add_job("Alice job", "jd", {})
        a.add_report("audit", "Alice audit", "<p>a</p>", {})
        self.assertEqual(b.list_jobs(), [])
        self.assertEqual(b.list_reports(), [])
        job_id = a.list_jobs()[0]["id"]
        self.assertIsNone(b.get_job(job_id))
        self.assertFalse(b.delete_job(job_id))
        self.assertIsNone(b.get_report_html(a.list_reports()[0]["id"]))
        # a brand-new workspace object (a new session) sees the stored data
        self.assertEqual(app.DbWorkspace(self.db, self.alice["id"], {}).list_jobs()[0]["title"], "Alice job")

    def test_db_workspace_caches_lists_until_something_changes(self):
        state = {}
        ws = app.DbWorkspace(self.db, self.alice["id"], state)
        ws.add_job("one", "jd", {})
        self.assertEqual(len(ws.list_jobs()), 1)
        auth.save_job(self.db, self.alice["id"], "added behind its back", "jd", {})
        self.assertEqual(len(ws.list_jobs()), 1)  # cached
        ws.add_job("two", "jd", {})  # a change invalidates the cache
        self.assertEqual(len(ws.list_jobs()), 3)

    def test_labels_tell_the_user_where_data_lives(self):
        self.assertIn("browser tab", app.SessionWorkspace({}).label)
        self.assertIn("wiped", app.DbWorkspace(self.db, 1, {}).label)  # sqlite: temporary


class AppSecretsTests(unittest.TestCase):
    class Fake:
        def __init__(self, secrets):
            self.secrets = secrets

    def test_reads_all_three_settings(self):
        smtp, admin, url = app.read_app_secrets(
            self.Fake({"smtp": {"host": "h", "port": 465}, "admin_email": " a@b.co ", "database_url": " postgresql://u:p@h/db "})
        )
        self.assertEqual((smtp["host"], admin, url), ("h", "a@b.co", "postgresql://u:p@h/db"))

    def test_missing_or_bad_values(self):
        self.assertEqual(app.read_app_secrets(self.Fake({})), (None, None, None))
        self.assertEqual(app.read_app_secrets(self.Fake({"admin_email": 5, "database_url": ""})), (None, None, None))

        class Exploding:
            @property
            def secrets(self):
                raise RuntimeError("no secrets file")

        self.assertEqual(app.read_app_secrets(Exploding()), (None, None, None))

    def test_apply_revised_jd_and_sign_out_keep_mode(self):
        state = {"jd_text": "old"}
        app._apply_revised_jd(state, "new")
        self.assertEqual(state["jd_text"], "new")
        state = {"user": 1, "dark_mode": True, "intro_done": True, "auth_mode": "Real account", "docs": [1]}
        app.sign_out(state)
        self.assertEqual(state, {"dark_mode": True, "intro_done": True, "auth_mode": "Real account"})

    def test_demo_user_is_marked_as_prototype(self):
        self.assertEqual(app.make_demo_user("a@b.co", "recruiter")["mode"], "prototype")


if __name__ == "__main__":
    unittest.main()
